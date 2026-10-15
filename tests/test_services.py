import asyncio
from io import BytesIO
import pytest
from backend.app.ai.providers import Completion, ProviderUnavailable
from backend.app.config import Settings
from backend.app.schemas import TranslationOutput, TranslationRequest
from backend.app.services.media import extract_document, export_docx
from backend.app.services.translation import segment, TranslationService


def test_segmentation_lossless():
    for text in ['Hello\n\nWorld', 'a' * 9000, ('A sentence. ' * 1000) + '\n終わり。', 'Hi\n  indented']:
        parts = segment(text)
        assert ''.join(parts) == text
        assert all(len(p) <= 1200 or p.isspace() for p in parts)


@pytest.mark.asyncio
async def test_retry_and_failover():
    class Broken:
        name = 'gemini'
        async def translate(self, *args):
            raise RuntimeError('sensitive provider error')
    class Working:
        name = 'deep-translator'
        supports_tone = False
        async def translate(self, *args):
            return Completion(TranslationOutput(translated_text='Bonjour', detected_language='en'))
    settings = Settings(_env_file=None, provider_order='gemini,deep-translator', provider_retries=1)
    service = TranslationService(settings, {'gemini': Broken(), 'deep-translator': Working()})
    result = await service.translate(TranslationRequest(text='Hello', target_lang='fr'))
    assert result['fallback'] is True and result['provider_failures'] == 2
    assert result['estimated_cost_usd'] is None and result['warnings']
    with pytest.raises(ProviderUnavailable):
        await service.translate(TranslationRequest(text='Hello', provider='gemini', allow_fallback=False))


@pytest.mark.asyncio
async def test_timeout_bounded():
    class Slow:
        name = 'gemini'
        async def translate(self, *args):
            await asyncio.sleep(0.1)
            raise RuntimeError('done')
    settings = Settings(_env_file=None)
    settings.provider_timeout = 0.01
    service = TranslationService(settings, {'gemini': Slow()})
    with pytest.raises(ProviderUnavailable):
        await service.translate(TranslationRequest(text='Hello'))
    assert len(service.pending) == 1
    await asyncio.sleep(0.15)
    assert not service.pending


def test_document_roundtrip():
    assert extract_document('नमस्ते\n\nWorld'.encode(), '.txt') == 'नमस्ते\n\nWorld'
    content = export_docx('नमस्ते\n\nWorld')
    assert extract_document(content, '.docx') == 'नमस्ते\n\nWorld'
    from docx import Document
    document = Document()
    document.add_paragraph('Before')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Left'
    table.cell(0, 1).text = 'Right'
    document.add_paragraph('After')
    stream = BytesIO()
    document.save(stream)
    assert extract_document(stream.getvalue(), '.docx') == 'Before\nLeft\tRight\nAfter'


@pytest.mark.asyncio
async def test_segment_boundaries_keep_spaces():
    class Echo:
        name = 'gemini'
        supports_tone = True
        async def translate(self, text, *args):
            return Completion(TranslationOutput(translated_text=text.strip(), detected_language='en'))
    text = ('A short sentence. ' * 120) + '\n\nDone'
    service = TranslationService(Settings(_env_file=None), {'gemini': Echo()})
    result = await service.translate(TranslationRequest(text=text))
    assert result['translated_text'] == text


def test_deep_translator_deadline(monkeypatch):
    import subprocess
    from backend.app.ai.providers import DeepTranslatorProvider
    def run(*args, **kwargs):
        assert kwargs['timeout'] == 2
        assert 'Hello' in kwargs['input']
        raise subprocess.TimeoutExpired(args[0], 2)
    monkeypatch.setattr(subprocess, 'run', run)
    with pytest.raises(subprocess.TimeoutExpired):
        DeepTranslatorProvider(timeout=2)._translate('Hello', 'en', 'fr')


def test_production_configuration():
    with pytest.raises(ValueError):
        Settings(_env_file=None, environment='production')
    with pytest.raises(ValueError):
        Settings(_env_file=None, environment='production', jwt_secret='x' * 40, database_url='sqlite:///test.db')
    settings = Settings(_env_file=None, environment='production', jwt_secret='x' * 40,
                        database_url='postgres://user:password@db/test')
    assert settings.database_url == 'postgresql+psycopg://user:password@db/test'
